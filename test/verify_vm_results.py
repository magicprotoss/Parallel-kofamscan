"""Stream-audit FAA workflow outputs against input FASTA and optional raw hits.

Run on the VM after the workflow, inside tmux. The input config is the workflow's
snakemake_config.yaml. Top-hit identities are independently checked for a small
sample subset; metadata and hit-to-sequence integrity are checked for all rows.
"""

import argparse
import csv
import gzip
import hashlib
import json
from pathlib import Path
import sqlite3
import tempfile

from openpyxl import load_workbook
import pyarrow.parquet as pq
import yaml


def records(path):
    opener = gzip.open if path.suffix == ".gz" else open
    with opener(path, "rt", encoding="utf-8") as handle:
        header, sequence = None, []
        for line in handle:
            if line.startswith(">"):
                if header is not None:
                    yield header, hashlib.sha256("".join(sequence).upper().encode("ascii")).hexdigest()
                header, sequence = line[1:].rstrip("\r\n"), []
            else:
                sequence.extend(line.split())
        if header is not None:
            yield header, hashlib.sha256("".join(sequence).upper().encode("ascii")).hexdigest()


def parquet_rows(path):
    with pq.ParquetFile(path) as parquet:
        for batch in parquet.iter_batches(batch_size=8192, use_threads=False):
            yield from batch.to_pylist()


def verify(config_path, report_path):
    config = yaml.safe_load(config_path.read_text())
    output = Path(config["output_path"]) / "final_results"
    inputs = {}
    for name in config["input_path"]:
        path = Path(name)
        sample = path.name.removesuffix(".gz")
        if not sample.endswith(".faa"):
            raise ValueError("This audit expects FAA inputs; FNA translation needs a separate check")
        inputs[sample[:-4]] = path
    samples = sorted(inputs)
    chosen = set(samples[:3] + samples[-2:])
    selected_records = {sample: [] for sample in chosen}
    selected_hashes = set()
    metadata_count = 0
    metadata = iter(parquet_rows(output / "sequence-metadata-map.parquet"))
    with tempfile.TemporaryDirectory(prefix="audit-") as scratch:
        index = sqlite3.connect(Path(scratch) / "hashes.sqlite")
        try:
            index.execute("PRAGMA cache_size = -8192")
            index.execute("CREATE TABLE hashes (hash TEXT PRIMARY KEY) WITHOUT ROWID")
            pending = []
            for sample in samples:
                for header, sequence_hash in records(inputs[sample]):
                    actual = next(metadata, None)
                    expected = {"sample_id": sample, "sequence_header": header, "sequence_hash": sequence_hash}
                    if actual != expected:
                        raise AssertionError(f"Metadata mismatch at occurrence {metadata_count + 1}: {sample}")
                    metadata_count += 1
                    pending.append((sequence_hash,))
                    if sample in chosen:
                        selected_records[sample].append((header, sequence_hash))
                        selected_hashes.add(sequence_hash)
                    if len(pending) == 8192:
                        with index:
                            index.executemany("INSERT OR IGNORE INTO hashes VALUES (?)", pending)
                        pending.clear()
            with index:
                index.executemany("INSERT OR IGNORE INTO hashes VALUES (?)", pending)
            if next(metadata, None) is not None:
                raise AssertionError("Extra metadata rows not present in input files")
            unique_count = index.execute("SELECT COUNT(*) FROM hashes").fetchone()[0]
            best = {}
            hit_count = 0
            with pq.ParquetFile(output / "hmm-hits.parquet") as hits:
                for batch in hits.iter_batches(batch_size=8192, use_threads=False):
                    rows = batch.to_pylist()
                    hashes = list({row["sequence_hash"] for row in rows})
                    for start in range(0, len(hashes), 500):
                        subset = hashes[start:start + 500]
                        query = "SELECT hash FROM hashes WHERE hash IN (" + ",".join("?" for _ in subset) + ")"
                        found = {row[0] for row in index.execute(query, subset)}
                        if found != set(subset):
                            raise AssertionError("HMM hit references a hash absent from sequence metadata")
                    for row in rows:
                        hit_count += 1
                        sequence_hash = row["sequence_hash"]
                        if (sequence_hash in selected_hashes
                                and row["E_value"] <= config["min_e_value"]
                                and row["hmmsearch_score"] >= config["min_score"]):
                            candidate = (row["E_value"], -row["hmmsearch_score"], row["KO"])
                            best[sequence_hash] = min(candidate, best.get(sequence_hash, candidate))
        finally:
            index.close()
    chunks = sorted((Path(config["working_dir"]) / "exec_annotation").glob("part_*/annotations.tsv"))
    raw_count = None
    if chunks:
        raw_count = 0
        raw_best = {}
        for chunk in chunks:
            with chunk.open(newline="") as handle:
                for row in csv.reader((line for line in handle if not line.startswith("#")), delimiter="\t"):
                    if row and len(row) == 7 and row[2]:
                        raw_count += 1
                        if (row[1] in selected_hashes and float(row[5]) <= config["min_e_value"]
                                and float(row[4]) >= config["min_score"]):
                            candidate = (float(row[5]), -float(row[4]), row[2])
                            raw_best[row[1]] = min(candidate, raw_best.get(row[1], candidate))
        if raw_count != hit_count:
            raise AssertionError(f"Raw hit count {raw_count} != Parquet hit count {hit_count}")
        if raw_best != best:
            raise AssertionError("Selected top hits differ between raw Kofam output and Parquet")
    workbook = load_workbook(output / "top-hits.xlsx", read_only=True)
    try:
        if "sample-id-map" in workbook.sheetnames:
            mapping_rows = list(workbook["sample-id-map"].values)
            if mapping_rows[0] != ("sample_id", "sheet_name"):
                raise AssertionError("Unexpected mapping columns")
            expected_mapping = [(sample, f"FAA{i}") for i, sample in enumerate(samples, 1)]
            if mapping_rows[1:] != expected_mapping:
                raise AssertionError("Sample-ID mapping differs from input manifest")
            mapping = dict(expected_mapping)
            expected_sheets = ["sample-id-map"] + list(mapping.values())
        else:
            mapping = {sample: sample for sample in samples}
            expected_sheets = samples
        if workbook.sheetnames != expected_sheets:
            raise AssertionError("Workbook sample sheets differ from input manifest")
        for sample in chosen:
            expected = [("unigene_id", "KO")] + [
                (header, best[sequence_hash][2]) for header, sequence_hash in selected_records[sample]
                if sequence_hash in best
            ]
            if list(workbook[mapping[sample]].values) != expected:
                raise AssertionError(f"Top hits differ for {sample}")
    finally:
        workbook.close()
    report = {"status": "passed", "samples": len(samples), "metadata_rows": metadata_count,
              "unique_sequence_hashes": unique_count, "hmm_hit_rows": hit_count,
              "raw_hit_rows": raw_count, "fully_checked_excel_samples": sorted(chosen)}
    report_path.write_text(json.dumps(report, indent=2) + "\n")
    print(json.dumps(report, indent=2), flush=True)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("config", type=Path)
    parser.add_argument("report", type=Path)
    args = parser.parse_args()
    verify(args.config, args.report)
