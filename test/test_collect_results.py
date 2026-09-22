"""Regression tests prepared for execution on the development VM."""

import csv
import hashlib
import importlib.util
from pathlib import Path
import shutil
from subprocess import CompletedProcess
import tempfile
import unittest
from unittest.mock import patch

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import load_workbook


REPO = Path(__file__).resolve().parents[1]
SCRIPTS = REPO / "parallel_kofam_scan" / "workflow" / "scripts"


def load_script(name):
    spec = importlib.util.spec_from_file_location(name, SCRIPTS / f"{name}.py")
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


collector = load_script("collect_results")
validation = load_script("validate_input")
splitting = load_script("concat_shuffle_split")


def digest(sequence):
    return hashlib.sha256(sequence.encode("ascii")).hexdigest()


class CollectionTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="kofam test's ")
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.output = self.root / "results"

    def sample(self, name, records, input_type="FAA"):
        directory = self.root / "validate_input" / name
        directory.mkdir(parents=True)
        schema = pa.schema([
            ("sequence_header", pa.string()), ("sequence_hash", pa.string()),
        ], metadata={b"hash_algorithm": b"sha256", b"input_type": input_type.encode()})
        pq.write_table(pa.Table.from_pylist([
            {"sequence_header": header, "sequence_hash": sequence_hash}
            for header, sequence_hash in records
        ], schema=schema), directory / "header_map.parquet", row_group_size=2)
        return directory / "header_map.parquet"

    def annotations(self, rows, part=1):
        path = self.root / "exec_annotation" / f"part_{part:03}" / "annotations.tsv"
        path.parent.mkdir(parents=True)
        with path.open("w", encoding="utf-8", newline="") as handle:
            handle.write('#\tgene name\tKO\tthrshld\tscore\tE-value\t"KO definition"\n')
            csv.writer(handle, delimiter="\t", lineterminator="\n").writerows(rows)
        return path

    def collect(self, **kwargs):
        # Small batches exercise the transitions between input files/row groups.
        with patch.object(collector, "BATCH_SIZE", 2), patch.object(collector, "LOOKUP_BATCH_SIZE", 1):
            collector.collect_results(self.root, self.output, **kwargs)

    def workbook(self):
        workbook = load_workbook(self.output / "top-hits.xlsx", read_only=True)
        self.addCleanup(workbook.close)
        return workbook

    def test_hits_are_not_multiplied_and_all_occurrences_survive(self):
        shared, other, absent = digest("MKT*"), digest("MKL*"), digest("MQQ*")
        header = '=gene 1 # full "description"\tretained'
        self.sample("sample-a", [(header, shared), ("copy", shared), ("no hit", absent)])
        self.sample("sample-b", [("shared b", shared), ("other", other)])
        self.sample("sample-empty", [])
        self.sample("sample-no-hits", [("no hit", absent)])
        definition = 'enzyme "quoted" #1\twith a tab'
        rows = [
            ["*", shared, "K00003", 90, 120, "1e-10", definition],
            ["", shared, "K00002", "", 130, "1e-10", "better score"],
            ["*", shared, "K00001", 90, 99, "0", "fails score cutoff"],
            ["", other, "K00004", "", 100, "0.001", "boundary"],
        ]
        self.annotations(rows[:2])
        self.annotations(rows[2:], part=2)
        self.collect()
        hits = pq.read_table(self.output / "hmm-hits.parquet")
        self.assertEqual(hits.num_rows, 4)
        self.assertEqual(hits.column_names[0], "sequence_hash")
        self.assertNotIn("sample_id", hits.column_names)
        self.assertEqual(hits.to_pylist()[0]["KO_definition"], definition)
        self.assertIsNone(hits.to_pylist()[1]["threshold"])
        metadata = pq.read_table(self.output / "sequence-metadata-map.parquet")
        self.assertEqual(metadata.column_names, ["sample_id", "sequence_header", "sequence_hash"])
        self.assertEqual(metadata.num_rows, 6)
        self.assertIn({"sample_id": "sample-a", "sequence_header": header,
                       "sequence_hash": shared}, metadata.to_pylist())
        workbook = self.workbook()
        self.assertEqual(workbook.sheetnames, ["sample-a", "sample-b", "sample-empty", "sample-no-hits"])
        self.assertEqual(list(workbook["sample-a"].values), [
            ("unigene_id", "KO"), (header, "K00002"), ("copy", "K00002"),
        ])
        self.assertEqual(workbook["sample-a"]["A2"].data_type, "s")
        self.assertEqual(list(workbook["sample-b"].values)[1:], [
            ("shared b", "K00002"), ("other", "K00004"),
        ])
        for name in ("sample-empty", "sample-no-hits"):
            self.assertEqual(list(workbook[name].values), [("unigene_id", "KO")])
        self.assertFalse((self.output / "annotations.parquet").exists())
        self.assertEqual(list(self.output.glob(".collect-*")), [])

    def test_custom_cutoffs_and_exact_ties_across_chunks(self):
        sequence_hash = digest("MKT*")
        self.sample("sample", [("gene", sequence_hash)])
        self.annotations([
            ["", sequence_hash, "K00009", "", 150, "1e-6", "eligible boundary"],
            ["*", sequence_hash, "K00001", 90, 149, "0", "fails custom score"],
            ["*", sequence_hash, "K00002", 90, 300, "1e-4", "fails custom E-value"],
        ])
        self.annotations([["", sequence_hash, "K00008", "", 150, "1e-6", "exact tie"]], part=2)
        self.collect(min_e_value=1e-6, min_score=150)
        self.assertEqual(list(self.workbook()["sample"].values)[1:], [("gene", "K00008")])
        self.assertEqual(pq.read_metadata(self.output / "hmm-hits.parquet").num_rows, 4)

    def test_long_name_maps_every_sample_with_input_type(self):
        sequence_hash = digest("MKT*")
        long_name = "b" * 32
        self.sample("a-short", [("gene", sequence_hash)])
        self.sample(long_name, [("gene", sequence_hash)], input_type="FNA")
        self.annotations([])
        self.collect()
        workbook = self.workbook()
        self.assertEqual(workbook.sheetnames, ["sample-id-map", "FAA1", "FNA2"])
        self.assertEqual(list(workbook["sample-id-map"].values), [
            ("sample_id", "sheet_name"), ("a-short", "FAA1"), (long_name, "FNA2"),
        ])
        self.assertEqual(pq.read_metadata(self.output / "sequence-metadata-map.parquet").num_rows, 2)

    def test_31_character_name_is_preserved(self):
        name = "a" * 31
        path = self.sample(name, [])
        sheets, needs_map = collector.sample_sheets([path])
        self.assertFalse(needs_map)
        self.assertEqual(sheets[0][1], name)

    def test_invalid_reserved_and_case_colliding_names(self):
        # Separate parents also let this test run on case-insensitive filesystems.
        for number, names in enumerate([
            ["bad[name]"], ["bad:name"], ["'quoted"], ["quoted'"],
            ["sample-id-map"], ["History"], ["Sample", "sample"],
        ]):
            paths = []
            for position, name in enumerate(names):
                directory = self.root / f"case-{number}-{position}" / name
                directory.mkdir(parents=True)
                source = self.sample(f"source-{number}-{position}", [])
                destination = directory / "header_map.parquet"
                shutil.copyfile(source, destination)
                paths.append(destination)
            with self.subTest(names=names):
                sheets, needs_map = collector.sample_sheets(paths)
                self.assertTrue(needs_map)
                self.assertEqual([sheet[1] for sheet in sheets],
                                 [f"FAA{i}" for i in range(1, len(names) + 1)])

    def test_empty_annotation_file_and_unannotated_rows(self):
        sequence_hash = digest("MKT*")
        self.sample("sample", [("gene", sequence_hash)])
        self.annotations([["", sequence_hash, "", "", "", "", ""]])
        empty = self.annotations([], part=2)
        empty.write_text("")
        self.collect()
        hits = pq.read_table(self.output / "hmm-hits.parquet")
        self.assertEqual(hits.num_rows, 0)
        self.assertEqual(hits.schema, collector.HIT_SCHEMA)
        self.assertEqual(list(self.workbook()["sample"].values), [("unigene_id", "KO")])

    def test_duplicate_full_headers_remain_distinct_occurrences(self):
        a, b = digest("MKT*"), digest("MKL*")
        self.sample("sample", [("same header", a), ("same header", b)])
        self.annotations([
            ["*", a, "K00001", 90, 120, "1e-5", "first protein"],
            ["*", b, "K00002", 90, 120, "1e-5", "second protein"],
        ])
        self.collect()
        self.assertEqual(list(self.workbook()["sample"].values)[1:], [
            ("same header", "K00001"), ("same header", "K00002"),
        ])

    def test_malformed_hits_do_not_publish_partial_results(self):
        self.sample("sample", [("gene", digest("MKT*"))])
        path = self.annotations([["*", "old-uuid", "K00001", 90, 100, "1e-4", "bad ID"]])
        with self.assertRaisesRegex(ValueError, "SHA-256 query ID"):
            self.collect()
        self.assertEqual(list(self.output.iterdir()), [])
        path.write_text("*\ttoo\tfew\tcolumns\n")
        with self.assertRaisesRegex(ValueError, "annotations.tsv:1"):
            self.collect()
        self.assertEqual(list(self.output.iterdir()), [])

    def test_legacy_uuid_maps_require_reprocessing(self):
        path = self.sample("sample", [])
        pq.write_table(pa.table({"id": ["gene"], "uuid": ["random-id"]}), path)
        self.annotations([])
        with self.assertRaisesRegex(ValueError, "legacy UUID maps"):
            self.collect()

    def test_failed_excel_export_removes_worksheet_temporary_files(self):
        sequence_hash = digest("MKT*")
        self.sample("a-finished", [("gene", sequence_hash)])
        self.sample("b-fails", [("x" * 32768, sequence_hash)])
        self.annotations([["*", sequence_hash, "K00001", 90, 100, "1e-5", "hit"]])
        with patch.object(tempfile, "tempdir", str(self.root)):
            with self.assertRaisesRegex(ValueError, "Excel's cell limit"):
                self.collect()
        self.assertEqual(list(self.output.iterdir()), [])
        self.assertEqual(list(self.root.glob("openpyxl.*")), [])

    def test_stock_faa_metadata_and_collection(self):
        # Uses the stock FASTA but synthetic hits; full Kofam/DB validation is separate.
        fixture = REPO / "test" / "data" / "SM1510S13_test.faa"
        directory = self.root / "validate_input" / fixture.stem
        directory.mkdir(parents=True)
        shutil.copyfile(fixture, directory / "unigene.faa")
        source = []
        for record in fixture.read_text().split(">")[1:]:
            lines = record.splitlines()
            source.append((lines[0], digest("".join(lines[1:]).upper())))
        with patch.object(validation, "BATCH_SIZE", 2):
            validation.make_header_map_df(directory)
        source_hashes = sorted({sequence_hash for _, sequence_hash in source})
        self.annotations([["*", sequence_hash, "K00001", 90, 100, "1e-5", "synthetic"]
                          for sequence_hash in source_hashes])
        self.collect()
        metadata = pq.read_table(self.output / "sequence-metadata-map.parquet").to_pylist()
        self.assertEqual([(row["sequence_header"], row["sequence_hash"]) for row in metadata], source)
        self.assertEqual(pq.read_metadata(self.output / "hmm-hits.parquet").num_rows, len(source_hashes))
        self.assertEqual(list(self.workbook()[fixture.stem].values)[1:],
                         [(header, "K00001") for header, _ in source])

    def test_hashes_canonicalize_case_wrapping_but_preserve_stops(self):
        directory = self.root / "hashing"
        directory.mkdir()
        (directory / "unigene.faa").write_text(
            '>gene 1 # description\nmkt\naa*\n>gene 2\nMKTAA*\n>gene 3\nMKTAA\n'
        )
        validation.make_header_map_df(directory, input_type="FNA")
        rows = pq.read_table(directory / "header_map.parquet").to_pylist()
        self.assertEqual([row["sequence_hash"] for row in rows],
                         [digest("MKTAA*"), digest("MKTAA*"), digest("MKTAA")])
        self.assertEqual(rows[0]["sequence_header"], "gene 1 # description")
        self.assertEqual(pq.read_schema(directory / "header_map.parquet").metadata[b"input_type"], b"FNA")
        self.assertEqual((directory / "unigene_hash.faa").read_text(),
                         f'>{digest("MKTAA*")}\nMKTAA*\n>{digest("MKTAA*")}\nMKTAA*\n>{digest("MKTAA")}\nMKTAA\n')

    def test_explicit_gzipped_faa_is_normalized_as_protein(self):
        with patch.object(validation, "run", return_value=CompletedProcess([], 0, "", "")) as run:
            result = validation.validate_input_and_convert_fna_to_faa("nucleotide alphabet.faa.gz", self.root)
        self.assertEqual(result, "FAA")
        run.assert_called_once()
        command = run.call_args.args[0]
        self.assertIn("protein", command)
        self.assertIn("nucleotide alphabet.faa.gz", command)
        self.assertIn(str(self.root / "unigene.faa"), command)

    def test_fna_translation_records_input_type(self):
        with patch.object(validation, "run", return_value=CompletedProcess([], 0, "", "")) as run:
            result = validation.validate_input_and_convert_fna_to_faa("input.fna", self.root)
        self.assertEqual(result, "FNA")
        self.assertEqual(run.call_args_list[1].args[0][-5:],
                         ["-T", "11", "input.fna", "-o", str(self.root / "unigene.faa")])

    def test_shuffle_failure_does_not_mark_partial_split_as_complete(self):
        sample_path = self.sample("sample", [])
        (sample_path.parent / "unigene_hash.faa").write_text(f'>{digest("MKT*")}\nMKT*\n')
        split_dir = self.root / "split"
        split_dir.mkdir()
        results = [CompletedProcess([], 0), CompletedProcess([], 0, "/bin/seqkit\n"),
                   CompletedProcess([], 0)]
        with patch.object(splitting, "run", side_effect=results), patch.object(splitting, "Popen") as popen:
            shuffle = popen.return_value.__enter__.return_value
            shuffle.wait.return_value = 1
            with self.assertRaisesRegex(SystemExit, "Failed to concatenate"):
                splitting.concat_shuffle_split(str(self.root / "validate_input"), 100,
                                               str(split_dir), 1)
        self.assertFalse((split_dir / "done").exists())


if __name__ == "__main__":
    unittest.main()
