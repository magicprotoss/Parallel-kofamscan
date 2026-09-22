#!/usr/bin/env python3
"""Collect normalized results without expanding hits across samples."""

import argparse
import csv
from contextlib import suppress
import math
import os
from pathlib import Path
import re
import sqlite3
import tempfile

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell


BATCH_SIZE = 8192
LOOKUP_BATCH_SIZE = 500  # Also works with SQLite's older 999-variable limit.
EXCEL_MAX_ROWS = 1048576
MAP_SHEET = "sample-id-map"
HASH_PATTERN = re.compile(r"[0-9a-f]{64}")
HIT_SCHEMA = pa.schema([
    ("sequence_hash", pa.string()),
    ("KO", pa.string()),
    ("threshold", pa.float64()),
    ("hmmsearch_score", pa.float64()),
    ("E_value", pa.float64()),
    ("KO_definition", pa.string()),
    ("if_score_higher_than_threshold", pa.string()),
], metadata={b"hash_algorithm": b"sha256"})
METADATA_SCHEMA = pa.schema([
    ("sample_id", pa.string()),
    ("sequence_header", pa.string()),
    ("sequence_hash", pa.string()),
], metadata={b"hash_algorithm": b"sha256"})


def parse_args():
    parser = argparse.ArgumentParser(
        prog="Parallel-kofamscan",
        description="Collect sequence metadata, HMM hits and per-sample top hits",
    )
    parser.add_argument("-wd", "--working_dir", required=True)
    parser.add_argument("-o", "--output_dir", required=True)
    parser.add_argument("-t", "--threads", default=1, type=int,
                        help="Retained for CLI compatibility; collection streams in a single worker")
    parser.add_argument("-e", "--min_e_value", default=0.001, type=float,
                        help="Maximum E-value retained during top-hit filtering (default: 0.001)")
    parser.add_argument("-s", "--min_score", default=100, type=int,
                        help="Minimum score retained during top-hit filtering (default: 100)")
    return parser.parse_args()


def sample_sheets(header_maps):
    """Preflight all titles; use deterministic aliases if any title is unsafe."""
    samples = []
    seen = {MAP_SHEET.casefold(), "history"}
    needs_map = False
    for path in header_maps:
        sample_id = path.parent.name
        schema = pq.read_schema(path)
        if schema.names != ["sequence_header", "sequence_hash"]:
            raise ValueError(
                f"{path}: expected sequence_header/sequence_hash; legacy UUID maps "
                "are incompatible. Rerun input validation and annotation."
            )
        metadata = schema.metadata or {}
        if metadata.get(b"hash_algorithm") != b"sha256":
            raise ValueError(f"{path}: expected SHA-256 sequence hashes; rerun input validation")
        input_type = metadata.get(b"input_type", b"FAA").decode("ascii")
        if input_type not in {"FAA", "FNA"}:
            raise ValueError(f"{path}: unknown input type {input_type!r}")
        if (not sample_id or len(sample_id) > 31
                or re.search(r"[\\/*?:\[\]\x00-\x1f]", sample_id)
                or sample_id.startswith("'") or sample_id.endswith("'")
                or sample_id.casefold() in seen):
            needs_map = True
        seen.add(sample_id.casefold())
        samples.append((sample_id, input_type, path))
    return [
        (sample_id, f"{input_type}{index}" if needs_map else sample_id, path)
        for index, (sample_id, input_type, path) in enumerate(samples, 1)
    ], needs_map


def annotation_batches(paths):
    """Read Kofam detail-tsv, including quoted definitions and empty thresholds."""
    rows = []
    for path in paths:
        with path.open(encoding="utf-8", newline="") as handle:
            for line_number, line in enumerate(handle, 1):
                if not line.strip() or line.startswith("#"):
                    continue
                try:
                    fields = next(csv.reader([line], delimiter="\t", strict=True))
                    marker, sequence_hash, ko, threshold, score, e_value, definition = fields
                    if not HASH_PATTERN.fullmatch(sequence_hash):
                        raise ValueError("expected a SHA-256 query ID; rerun validation and annotation")
                    # Optional --report-unannotated lines are not HMM hits.
                    if not ko and not any(fields[2:]) and not marker:
                        continue
                    if marker not in {"", "*"} or not ko:
                        raise ValueError("invalid hit marker or missing KO")
                    threshold = None if threshold in {"", "-"} else float(threshold)
                    score, e_value = float(score), float(e_value)
                    if (not math.isfinite(score) or not math.isfinite(e_value)
                            or e_value < 0 or (threshold is not None and not math.isfinite(threshold))):
                        raise ValueError("invalid numeric hit value")
                except (ValueError, csv.Error) as error:
                    raise ValueError(f"{path}:{line_number}: {error}") from error
                rows.append(dict(zip(HIT_SCHEMA.names, (
                    sequence_hash, ko, threshold, score, e_value, definition, marker,
                ))))
                if len(rows) == BATCH_SIZE:
                    yield rows
                    rows = []
    if rows:
        yield rows


def write_hits(paths, output_path, index, min_e_value, min_score):
    # Only one qualifying hit per distinct hash is kept in this disk-backed index.
    # Rank by E-value, then score; KO breaks exact ties reproducibly across chunks.
    index.execute("""
        CREATE TABLE top_hits (
            sequence_hash TEXT PRIMARY KEY, KO TEXT NOT NULL,
            E_value REAL NOT NULL, hmmsearch_score REAL NOT NULL
        ) WITHOUT ROWID
    """)
    upsert = """
        INSERT INTO top_hits VALUES (?, ?, ?, ?)
        ON CONFLICT(sequence_hash) DO UPDATE SET
            KO = excluded.KO, E_value = excluded.E_value,
            hmmsearch_score = excluded.hmmsearch_score
        WHERE excluded.E_value < top_hits.E_value
           OR (excluded.E_value = top_hits.E_value
               AND excluded.hmmsearch_score > top_hits.hmmsearch_score)
           OR (excluded.E_value = top_hits.E_value
               AND excluded.hmmsearch_score = top_hits.hmmsearch_score
               AND excluded.KO < top_hits.KO)
    """
    count = 0
    with pq.ParquetWriter(output_path, HIT_SCHEMA, compression="zstd") as writer:
        for rows in annotation_batches(paths):
            writer.write_table(pa.Table.from_pylist(rows, schema=HIT_SCHEMA))
            # Reduce repeated hits within this bounded batch before touching disk.
            best = {}
            for row in rows:
                if row["E_value"] > min_e_value or row["hmmsearch_score"] < min_score:
                    continue
                rank = (row["E_value"], -row["hmmsearch_score"], row["KO"])
                previous = best.get(row["sequence_hash"])
                if previous is None or rank < previous:
                    best[row["sequence_hash"]] = rank
            with index:
                index.executemany(upsert, (
                    (sequence_hash, ko, e_value, -negative_score)
                    for sequence_hash, (e_value, negative_score, ko) in best.items()
                ))
            count += len(rows)
    return count


def lookup_top_hits(index, hashes):
    hashes = list(set(hashes))
    hits = {}
    for start in range(0, len(hashes), LOOKUP_BATCH_SIZE):
        batch = hashes[start:start + LOOKUP_BATCH_SIZE]
        placeholders = ",".join("?" for _ in batch)
        hits.update(index.execute(
            f"SELECT sequence_hash, KO FROM top_hits WHERE sequence_hash IN ({placeholders})",
            batch,
        ))
    return hits


def append_text(sheet, values):
    # Preserve FASTA headers/sample names literally, including leading '='.
    cells = []
    for value in values:
        if len(value) > 32767:
            raise ValueError(f"Text exceeds Excel's cell limit in sheet {sheet.title!r}")
        cell = WriteOnlyCell(sheet, value=value)
        cell.data_type = "s"
        cells.append(cell)
    sheet.append(cells)


def write_metadata_and_excel(samples, needs_map, output_dir, index):
    workbook = Workbook(write_only=True)
    count = 0
    try:
        if needs_map:
            sheet = workbook.create_sheet(MAP_SHEET)
            append_text(sheet, ("sample_id", "sheet_name"))
            for sample_id, title, _ in samples:
                append_text(sheet, (sample_id, title))
            sheet.close()
        with pq.ParquetWriter(output_dir / "sequence-metadata-map.parquet",
                              METADATA_SCHEMA, compression="zstd") as writer:
            for sample_id, title, path in samples:
                sheet = workbook.create_sheet(title)
                append_text(sheet, ("unigene_id", "KO"))
                sheet_rows = 1
                with pq.ParquetFile(path) as header_map:
                    for batch in header_map.iter_batches(batch_size=BATCH_SIZE, use_threads=False):
                        columns = batch.to_pydict()
                        headers = columns["sequence_header"]
                        hashes = columns["sequence_hash"]
                        if any(h is None or not HASH_PATTERN.fullmatch(h) for h in hashes):
                            raise ValueError(f"{path}: invalid sequence hash")
                        if any(h is None or not h for h in headers):
                            raise ValueError(f"{path}: empty sequence header")
                        writer.write_table(pa.Table.from_pydict({
                            "sample_id": [sample_id] * len(hashes),
                            "sequence_header": headers,
                            "sequence_hash": hashes,
                        }, schema=METADATA_SCHEMA))
                        hits = lookup_top_hits(index, hashes)
                        for header, sequence_hash in zip(headers, hashes):
                            if sequence_hash in hits:
                                if sheet_rows >= EXCEL_MAX_ROWS:
                                    raise ValueError(f"Sample {sample_id!r} exceeds Excel's worksheet row limit")
                                append_text(sheet, (header, hits[sequence_hash]))
                                sheet_rows += 1
                        count += len(hashes)
                # Release the worksheet's temporary-file handle before the next sample.
                sheet.close()
        workbook.save(output_dir / "top-hits.xlsx")
    except BaseException:
        for sheet in workbook.worksheets:
            # Public close() finishes XML but does not remove its temporary file.
            # Openpyxl normally does that in save(); failed exports need cleanup
            # here. Try every sheet without masking the original export error.
            with suppress(Exception):
                if not sheet.closed:
                    sheet.close()
            sheet_writer = getattr(sheet, "_writer", None)
            if sheet_writer is not None:
                with suppress(Exception):
                    sheet_writer.cleanup()
        raise
    finally:
        workbook.close()
    return count


def collect_results(working_dir, output_dir, threads=1, min_e_value=0.001, min_score=100):
    if not math.isfinite(min_e_value) or min_e_value < 0 or not math.isfinite(min_score):
        raise ValueError("Top-hit thresholds must be finite, with E-value >= 0")
    working_dir, output_dir = Path(working_dir), Path(output_dir)
    header_maps = sorted((working_dir / "validate_input").glob("*/header_map.parquet"))
    annotations = sorted((working_dir / "exec_annotation").glob("part_*/annotations.tsv"))
    if not header_maps:
        raise FileNotFoundError(f"No header maps found under {working_dir / 'validate_input'}")
    if not annotations:
        raise FileNotFoundError(f"No annotation chunks found under {working_dir / 'exec_annotation'}")
    samples, needs_map = sample_sheets(header_maps)
    output_dir.mkdir(parents=True, exist_ok=True)
    # Stage all outputs so a parsing/export failure does not publish partial results.
    with tempfile.TemporaryDirectory(prefix=".collect-", dir=output_dir) as staging:
        staging = Path(staging)
        index = sqlite3.connect(staging / "top-hits.sqlite")
        try:
            index.execute("PRAGMA cache_size = -8192")  # 8 MiB page cache.
            index.execute("PRAGMA temp_store = FILE")
            print("Streaming HMM hits and indexing the best hit per sequence hash...", flush=True)
            hit_count = write_hits(annotations, staging / "hmm-hits.parquet", index,
                                   min_e_value, min_score)
            print("Streaming sequence metadata and per-sample Excel sheets...", flush=True)
            sequence_count = write_metadata_and_excel(samples, needs_map, staging, index)
        finally:
            index.close()
        for name in ("hmm-hits.parquet", "sequence-metadata-map.parquet", "top-hits.xlsx"):
            os.replace(staging / name, output_dir / name)
    print(f"Collected {hit_count} hits and {sequence_count} sequence occurrences "
          f"across {len(samples)} samples.", flush=True)


def main():
    args = parse_args()
    collect_results(args.working_dir, args.output_dir, args.threads,
                    args.min_e_value, args.min_score)


if __name__ == "__main__":
    main()
